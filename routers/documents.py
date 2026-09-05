import os
import io
import json
import httpx
import re
from fastapi import APIRouter, Depends, HTTPException, UploadFile, File, Form, BackgroundTasks, status
from sqlalchemy.orm import Session
from pypdf import PdfReader
import docx
from database import get_db, SessionLocal
from models import Document, ExtractedData, Log, User

router = APIRouter(prefix="/documents", tags=["Documents"])

POLZA_API_KEY = os.getenv("POLZA_API_KEY", "pza_FwMrKaXpRs-bs0ELsLmAyayjzdJnXkQH")
POLZA_API_URL = "https://api.polza.ai/v1/chat/completions"

UPLOAD_DIR = "uploads"
os.makedirs(UPLOAD_DIR, exist_ok=True)


# --- Вспомогательные функции ---

def extract_text(file_path: str) -> str:
    text = ""
    try:
        if file_path.lower().endswith('.pdf'):
            with open(file_path, "rb") as f:
                reader = PdfReader(f)
                for page in reader.pages:
                    text += page.extract_text() + "\n"
        elif file_path.lower().endswith('.docx'):
            doc = docx.Document(file_path)
            for para in doc.paragraphs:
                text += para.text + "\n"
    except Exception as e:
        print(f"Ошибка чтения файла {file_path}: {e}")
    return text

async def fetch_ai_attributes(text: str) -> dict:
    headers = {
        "Authorization": f"Bearer {POLZA_API_KEY}",
        "Content-Type": "application/json"
    }
    prompt = f"""
    Извлеки из текста договора 4 атрибута. 
    Верни СТРОГО JSON: {{"number": "номер", "date": "дата", "parties": "стороны", "amount": "сумма"}}. 
    Никаких вступительных слов, только JSON. Если чего-то нет, пиши "Не найдено". 
    Текст: {text[:4000]}
    """
    payload = {
        "model": "openai/gpt-4o-mini",
        "messages": [{"role": "user", "content": prompt}],
        "temperature": 0.1
    }
    
    async with httpx.AsyncClient(timeout=30.0) as client:
        try:
            res = await client.post(POLZA_API_URL, headers=headers, json=payload)
            res.raise_for_status()
            ai_text = res.json()['choices'][0]['message']['content'].strip()
            
            match = re.search(r'\{.*\}', ai_text, re.DOTALL)
            if match:
                return json.loads(match.group(0))
            else:
                return {"error": "ИИ не вернул валидный JSON"}
        except Exception as e:
            return {"error": f"Ошибка запроса к ИИ: {str(e)}"}


# --- Фоновая задача ---

async def process_document_background(document_id: int, file_path: str):
    """Теперь задача сама создает себе сессию БД и корректно ее закрывает."""
    db = SessionLocal() # Открываем свежую сессию
    try:
        doc = db.query(Document).filter(Document.id == document_id).first()
        if not doc:
            return

        doc.status = "processing"
        db.add(Log(document_id=doc.id, message="Начато извлечение текста"))
        db.commit()

        raw_text = extract_text(file_path)
        if not raw_text.strip():
            raise ValueError("Не удалось извлечь текст или файл пуст")

        db.add(Log(document_id=doc.id, message="Текст извлечен, отправка запроса к LLM"))
        db.commit()

        ai_data = await fetch_ai_attributes(raw_text)
        
        extracted = ExtractedData(document_id=doc.id, extracted_data=ai_data)
        db.add(extracted)
        
        doc.status = "completed"
        db.add(Log(document_id=doc.id, message="Анализ успешно завершен"))
        db.commit()

    except Exception as e:
        print(f"Критическая ошибка в фоне: {e}")
        doc.status = "failed"
        db.add(Log(document_id=document_id, message=f"Критическая ошибка: {str(e)}"))
        db.commit()
    finally:
        db.close() # Обязательно закрываем сессию!


# --- API Эндпоинты ---

@router.post("/upload", status_code=status.HTTP_202_ACCEPTED)
async def upload_document(
    background_tasks: BackgroundTasks,
    user_id: int = Form(...),
    file: UploadFile = File(...),
    db: Session = Depends(get_db)
):
    user = db.query(User).filter(User.id == user_id).first()
    if not user:
        raise HTTPException(status_code=404, detail="Пользователь не найден")

    if not file.filename.lower().endswith(('.pdf', '.docx')):
        raise HTTPException(status_code=400, detail="Разрешены только файлы .pdf и .docx")

    file_path = os.path.join(UPLOAD_DIR, file.filename)
    with open(file_path, "wb") as f:
        f.write(await file.read())

    new_doc = Document(
        user_id=user_id,
        filename=file.filename,
        file_path=file_path,
        status="pending"
    )
    db.add(new_doc)
    db.commit()
    db.refresh(new_doc)
    
    db.add(Log(document_id=new_doc.id, message="Файл загружен, ожидает обработки"))
    db.commit()

    background_tasks.add_task(process_document_background, new_doc.id, file_path)

    return {
        "status": "success", 
        "message": "Файл загружен и отправлен в очередь на обработку",
        "document_id": new_doc.id
    }

@router.get("/{user_id}")
def get_user_documents(user_id: int, db: Session = Depends(get_db)):
    docs = db.query(Document).filter(Document.user_id == user_id).order_by(Document.id.desc()).all()
    
    result = []
    for doc in docs:
        result.append({
            "id": doc.id,
            "filename": doc.filename,
            "status": doc.status,
            "uploaded_at": doc.uploaded_at,
            "extracted_data": doc.extracted_data.extracted_data if doc.extracted_data else None,
            "logs": [log.message for log in doc.logs]
        })
    return result